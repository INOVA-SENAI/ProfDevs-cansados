"""Prévia dos slides: desenha o .pptx gerado como PNG, sem PowerPoint nem LibreOffice.

Portado de `exportar_previa.py` do repositório icrcode/teste-kiro (licença MIT; ver
LICENSES/teste-kiro.txt). Desenha as formas que o gerador usa (retângulos, círculos,
polígonos, imagens e caixas de texto em Open Sans); outros recursos do PowerPoint não aparecem.
"""

import io
from functools import cache
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from senai_slides.tema import FONTES

EMU = 914400
ESCALA = 2  # desenha em 2x e reduz, para bordas e textos suaves
NS_A = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


@cache
def _fonte(bold: bool, italic: bool, px: int) -> ImageFont.FreeTypeFont:
    estilo = (
        "BoldItalic" if bold and italic else "Bold" if bold else "Italic" if italic else "Regular"
    )
    return ImageFont.truetype(str(FONTES / f"OpenSans-{estilo}.ttf"), max(1, px))


def _cor(fill) -> str | None:
    try:
        if fill.type == 1:  # MSO_FILL.SOLID
            return "#" + str(fill.fore_color.rgb)
    except (AttributeError, TypeError):
        pass
    return None


def _quebrar(draw, texto, fonte, largura) -> list[str]:
    linhas = []
    for paragrafo in texto.split("\n"):
        atual = ""
        for palavra in paragrafo.split(" "):
            teste = f"{atual} {palavra}".strip()
            if not atual or draw.textlength(teste, font=fonte) <= largura:
                atual = teste
            else:
                linhas.append(atual)
                atual = palavra
        linhas.append(atual)
    return linhas


def _texto(draw, forma, px) -> None:
    tf = forma.text_frame
    x, y, w, h = px(forma.left), px(forma.top), px(forma.width), px(forma.height)
    ml, mr = px(tf.margin_left or 0), px(tf.margin_right or 0)
    mt, mb = px(tf.margin_top or 0), px(tf.margin_bottom or 0)
    linhas = []
    for par in tf.paragraphs:
        if not par.runs:
            continue
        run = par.runs[0]
        tamanho = (run.font.size.pt if run.font.size else 18) / 72 * px(EMU)
        fonte = _fonte(bool(run.font.bold), bool(run.font.italic), int(tamanho))
        cor = "#" + str(run.font.color.rgb) if run.font.color and run.font.color.type else "#000000"
        espaco = par.line_spacing if isinstance(par.line_spacing, float) else 1.0
        texto = "".join(r.text for r in par.runs)
        for linha in _quebrar(draw, texto, fonte, w - ml - mr + 1):
            linhas.append((linha, fonte, cor, tamanho * 1.36 * espaco, par.alignment))
    altura = sum(linha[3] for linha in linhas)
    ancora = tf.vertical_anchor
    if ancora == MSO_ANCHOR.MIDDLE or (
        ancora is None and forma.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
    ):
        cy = y + (h - altura) / 2
    elif ancora == MSO_ANCHOR.BOTTOM:
        cy = y + h - mb - altura
    else:
        cy = y + mt
    for linha, fonte, cor, alt, alinhamento in linhas:
        largura = draw.textlength(linha, font=fonte)
        if alinhamento == PP_ALIGN.CENTER:
            cx = x + (w - largura) / 2
        elif alinhamento == PP_ALIGN.RIGHT:
            cx = x + w - mr - largura
        else:
            cx = x + ml
        draw.text((cx, cy + alt * 0.08), linha, font=fonte, fill=cor)
        cy += alt


def renderizar(prs, slide, largura: int = 1920) -> Image.Image:
    ppi = largura / (prs.slide_width / EMU) * ESCALA

    def px(v):
        return v / EMU * ppi

    w_total, h_total = round(px(prs.slide_width)), round(px(prs.slide_height))
    im = Image.new("RGB", (w_total, h_total), _cor(slide.background.fill) or "#FFFFFF")
    draw = ImageDraw.Draw(im)
    for forma in slide.shapes:
        x, y, w, h = px(forma.left), px(forma.top), px(forma.width), px(forma.height)
        tipo = forma.shape_type
        if tipo == MSO_SHAPE_TYPE.PICTURE:
            figura = Image.open(io.BytesIO(forma.image.blob)).convert("RGBA")
            figura = figura.resize((max(1, round(w)), max(1, round(h))), Image.Resampling.LANCZOS)
            im.paste(figura, (round(x), round(y)), figura)
        elif tipo == MSO_SHAPE_TYPE.FREEFORM:
            pontos = [
                (x + int(p.get("x")) / EMU * ppi, y + int(p.get("y")) / EMU * ppi)
                for p in forma._element.findall(".//a:pt", NS_A)
            ]
            cor = _cor(forma.fill)
            if pontos and cor:
                draw.polygon(pontos, fill=cor)
        elif tipo == MSO_SHAPE_TYPE.AUTO_SHAPE:
            cor = _cor(forma.fill)
            if cor:
                caixa = [x, y, x + w, y + h]
                estilo = forma.auto_shape_type
                if estilo == MSO_SHAPE.OVAL:
                    draw.ellipse(caixa, fill=cor)
                elif estilo == MSO_SHAPE.ROUNDED_RECTANGLE:
                    raio = forma.adjustments[0] * min(w, h)
                    draw.rounded_rectangle(caixa, radius=raio, fill=cor)
                else:
                    draw.rectangle(caixa, fill=cor)
            if forma.has_text_frame and forma.text_frame.text:
                _texto(draw, forma, px)
        if tipo == MSO_SHAPE_TYPE.TEXT_BOX:
            _texto(draw, forma, px)
    return im.resize((w_total // ESCALA, h_total // ESCALA), Image.Resampling.LANCZOS)


def exportar(prs, pasta: Path, largura: int = 1920) -> list[Path]:
    """Salva slide_01.png, slide_02.png... e <pasta>.pdf com todos os slides."""
    for antigo in pasta.glob("slide_*.png"):
        antigo.unlink()
    imagens = [renderizar(prs, slide, largura) for slide in prs.slides]
    caminhos = []
    for i, imagem in enumerate(imagens, start=1):
        caminho = pasta / f"slide_{i:02d}.png"
        imagem.save(caminho, optimize=True)
        caminhos.append(caminho)
    # 144 dpi deixa a página em 13,33 x 7,5 polegadas, o tamanho padrão de slide 16:9.
    imagens[0].save(
        pasta / f"{pasta.name}.pdf",
        save_all=True,
        append_images=imagens[1:],
        resolution=largura / 13.333,
    )
    return caminhos
