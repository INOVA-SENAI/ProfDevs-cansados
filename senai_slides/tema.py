"""Tema SENAI 2026 em PowerPoint: cores, fontes, ícones e as peças de desenho dos slides.

Portado do gerador de slides do repositório icrcode/teste-kiro (licença MIT, Ícaro Caldeira
Botelho; ver LICENSES/teste-kiro.txt), que segue o modelo oficial
`assets/modelo_senai_2026_identidade_nova.pptx`. Medidas em polegadas; o slide tem 13,333 x
7,5 pol (16:9).
"""

import io
import os
import shutil
import sys
from dataclasses import dataclass, field
from functools import cache
from pathlib import Path

import matplotlib
import numpy as np
from lxml import etree
from matplotlib.figure import Figure
from matplotlib.font_manager import fontManager
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

PASTA = Path(__file__).resolve().parent
ASSETS = PASTA / "assets"
FONTES = ASSETS / "fonts"
CACHE = ASSETS / ".cache"
LOGO_SENAI = ASSETS / "senai.png"
LOGO_CURSO = ASSETS / "tecnico.png"
SIMBOLOS_TTF = FONTES / "MaterialSymbolsRounded.ttf"
SIMBOLOS_CP = FONTES / "MaterialSymbolsRounded.codepoints"
OPEN_SANS = ["Regular", "Italic", "Light", "SemiBold", "Bold", "BoldItalic"]
FONTE = "Open Sans"


def rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa.lstrip("#"))


H_AZUL = "#164194"  # azul SENAI (títulos, painéis)
H_LARANJA = "#E84910"  # laranja SENAI (acentos)
H_BRANCO = "#FFFFFF"
H_TINTA = "#1E2A44"  # texto de dados
H_CINZA = "#5B6475"  # texto secundário
H_LINHA = "#DDE3EE"  # grades e divisores
H_SUPERF = "#F3F5FA"  # fundo dos cartões
H_AZUL_BG = "#E6ECF7"  # fundo dos ícones
H_LAR_BG = "#FDE9E1"
H_AZUL_TXT = "#C9D6F0"  # texto secundário sobre azul
# Séries dos gráficos, nesta ordem (validada para daltonismo).
H_SERIE_1, H_SERIE_2, H_SERIE_3 = "#7FA3E3", "#2458B8", H_LARANJA

AZUL, LARANJA, BRANCO = rgb(H_AZUL), rgb(H_LARANJA), rgb(H_BRANCO)
CINZA, SUPERF, LINHA = rgb(H_CINZA), rgb(H_SUPERF), rgb(H_LINHA)
AZUL_BG, LAR_BG = rgb(H_AZUL_BG), rgb(H_LAR_BG)
AZUL_TXT, TINTA = rgb(H_AZUL_TXT), rgb(H_TINTA)

W, H = 13.333, 7.5
X0, X1 = 0.9, 12.58  # margens do conteúdo
Y0, Y1 = 2.0, 6.75  # área útil abaixo do cabeçalho
DPI = 300  # resolução dos gráficos
I = Inches  # noqa: E741


# Recursos: fontes, ícones e logos


@cache
def registrar_fontes() -> None:
    """Deixa a Open Sans disponível para o matplotlib (gráficos)."""
    for estilo in OPEN_SANS:
        fontManager.addfont(str(FONTES / f"OpenSans-{estilo}.ttf"))


def instalar_fontes_usuario() -> list[str]:
    """Instala a Open Sans só para o usuário atual no Windows, sem administrador, para o
    PowerPoint mostrar a fonte certa. Devolve as fontes que foram instaladas."""
    if sys.platform != "win32":
        return []
    import winreg

    sistema = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    usuario = Path(os.environ["LOCALAPPDATA"]) / "Microsoft" / "Windows" / "Fonts"
    chave = r"Software\Microsoft\Windows NT\CurrentVersion\Fonts"
    novas = []
    for estilo in OPEN_SANS:
        nome = f"OpenSans-{estilo}.ttf"
        if (sistema / nome).exists() or (usuario / nome).exists():
            continue
        usuario.mkdir(parents=True, exist_ok=True)
        destino = usuario / nome
        shutil.copy2(FONTES / nome, destino)
        familia, variante = ImageFont.truetype(str(destino), 12).getname()
        rotulo = familia if variante == "Regular" else f"{familia} {variante}"
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, chave, 0, winreg.KEY_SET_VALUE) as k:
            winreg.SetValueEx(k, f"{rotulo} (TrueType)", 0, winreg.REG_SZ, str(destino))
        novas.append(nome)
    return novas


@cache
def codepoints() -> dict[str, str]:
    linhas = SIMBOLOS_CP.read_text(encoding="utf-8").splitlines()
    return dict(linha.split() for linha in linhas if linha.strip())


@cache
def _fonte_icone(tamanho: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(SIMBOLOS_TTF), tamanho)


def icone_png(nome: str, cor: str = H_AZUL, px: int = 384) -> str:
    """Desenha um Material Symbol como PNG transparente (em 2x e reduzido) e devolve o
    caminho do arquivo em cache."""
    caminho = CACHE / "icones" / f"{nome}_{cor.lstrip('#')}_{px}.png"
    if not caminho.exists():
        s = px * 2
        im = Image.new("RGBA", (s, s), (0, 0, 0, 0))
        glifo = chr(int(codepoints()[nome], 16))
        ImageDraw.Draw(im).text((s / 2, s / 2), glifo, font=_fonte_icone(s), fill=cor, anchor="mm")
        caminho.parent.mkdir(parents=True, exist_ok=True)
        im.resize((px, px), Image.Resampling.LANCZOS).save(caminho)
    return str(caminho)


def _em_cache(nome: str, criar) -> str:
    destino = CACHE / nome
    if not destino.exists():
        destino.parent.mkdir(parents=True, exist_ok=True)
        criar().save(destino)
    return str(destino)


def logo_curso() -> str:
    """Logo "Técnico DESI" sem a margem transparente."""

    def criar():
        im = Image.open(LOGO_CURSO).convert("RGBA")
        im = im.crop(im.getchannel("A").getbbox())
        im.thumbnail((1600, 1600), Image.Resampling.LANCZOS)
        return im

    return _em_cache("tecnico_recortado.png", criar)


def logo_branco() -> str:
    """Logo SENAI em branco, para os fundos azuis."""

    def criar():
        im = Image.open(LOGO_SENAI).convert("RGBA")
        branco = Image.new("RGBA", im.size, (255, 255, 255, 0))
        branco.putalpha(im.getchannel("A"))
        return branco

    return _em_cache("senai_branco.png", criar)


def fundo_azul() -> str:
    """Fundo azul SENAI das divisórias e do encerramento, com um degradê suave na diagonal.

    O modelo do teste-kiro usava uma colagem de fotos de alunos a 8%; ela não entra aqui,
    porque são fotos de pessoas (LGPD) de uma turma específica."""

    def criar():
        vertical = Image.linear_gradient("L").resize((1920, 1080))
        horizontal = Image.linear_gradient("L").rotate(90).resize((1920, 1080))
        diagonal = Image.blend(vertical, horizontal.transpose(Image.Transpose.FLIP_LEFT_RIGHT), 0.5)
        return ImageOps.colorize(diagonal, black="#1A4AA3", white="#0F3279")

    return _em_cache("fundo_azul.jpg", criar)


# Peças do PowerPoint


def set_bg(sl, cor) -> None:
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = cor


def rect(sl, x, y, w, h, fill=None, line=None, lw=1.0, raio=None, forma=None):
    tipo = forma or (MSO_SHAPE.ROUNDED_RECTANGLE if raio else MSO_SHAPE.RECTANGLE)
    s = sl.shapes.add_shape(tipo, I(x), I(y), I(w), I(h))
    if raio:
        s.adjustments[0] = raio
    if fill is not None:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line is not None:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def circulo(sl, x, y, d, fill):
    return rect(sl, x, y, d, d, fill=fill, forma=MSO_SHAPE.OVAL)


def poligono(sl, pontos, fill):
    construtor = sl.shapes.build_freeform(I(pontos[0][0]), I(pontos[0][1]), scale=1.0)
    construtor.add_line_segments([(I(px), I(py)) for px, py in pontos[1:]], close=True)
    s = construtor.convert_to_shape()
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


_ALINHAMENTO = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
_ANCORA = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def txt(
    sl,
    texto,
    x,
    y,
    w,
    h,
    size=16,
    bold=False,
    italic=False,
    cor=AZUL,
    align="l",
    anchor="t",
    espaco=1.1,
    depois=0,
):
    """Caixa de texto Open Sans sem margens internas; "\\n" separa parágrafos."""
    tb = sl.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCORA[anchor]
    for i, linha in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALINHAMENTO[align]
        p.line_spacing = espaco
        if depois:
            p.space_after = Pt(depois)
        r = p.add_run()
        r.text = linha
        r.font.name, r.font.size, r.font.bold, r.font.italic = FONTE, Pt(size), bold, italic
        r.font.color.rgb = cor
    return tb


def _arquivo_fonte(bold: bool, italic: bool) -> str:
    estilo = (
        "BoldItalic" if bold and italic else "Bold" if bold else "Italic" if italic else "Regular"
    )
    return str(FONTES / f"OpenSans-{estilo}.ttf")


ALTURA_DA_LINHA = 1.36  # altura de uma linha da Open Sans, em múltiplos do tamanho da fonte
TAMANHO_MINIMO = 7


@cache
def _fonte_de_medida(bold: bool, italic: bool) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(_arquivo_fonte(bold, italic), 100)


def _linhas(texto: str, fonte, largura: float) -> list[str]:
    linhas = []
    for paragrafo in texto.split("\n"):
        atual = ""
        for palavra in paragrafo.split():
            tentativa = f"{atual} {palavra}".strip()
            if not atual or fonte.getlength(tentativa) <= largura:
                atual = tentativa
            else:
                linhas.append(atual)
                atual = palavra
        linhas.append(atual)
    return linhas


def tamanho_que_cabe(texto, w, h, max_size, bold=False, italic=False, espaco=1.0) -> float:
    """Maior tamanho (em pontos, de 0,5 em 0,5) em que o texto cabe em w x h polegadas.

    Substitui o `fit_text` do python-pptx, que falha quando uma palavra sozinha não cabe na
    largura. A medida é a mesma da prévia: Open Sans, quebra por palavras."""
    fonte = _fonte_de_medida(bold, italic)
    tamanho = float(max_size)
    while tamanho > TAMANHO_MINIMO:
        largura = w * 72 * 100 / tamanho  # largura da caixa medida com a fonte de 100 px
        linhas = _linhas(texto, fonte, largura)
        altura = len(linhas) * tamanho * ALTURA_DA_LINHA * espaco / 72
        if altura <= h and all(fonte.getlength(linha) <= largura for linha in linhas):
            return tamanho
        tamanho -= 0.5
    return TAMANHO_MINIMO


def txt_fit(sl, texto, x, y, w, h, max_size, bold=False, italic=False, **kw):
    """Como `txt`, mas reduz a fonte até o texto caber na caixa."""
    kw.setdefault("espaco", 1.0)
    tamanho = tamanho_que_cabe(texto, w, h, max_size, bold, italic, kw["espaco"])
    return txt(sl, texto, x, y, w, h, size=tamanho, bold=bold, italic=italic, **kw)


def icone(sl, nome, x, y, tam, cor=H_AZUL):
    return sl.shapes.add_picture(icone_png(nome, cor), I(x), I(y), I(tam), I(tam))


def icone_badge(sl, nome, x, y, d, fundo=AZUL_BG, cor=H_AZUL):
    """Ícone dentro de um círculo suave."""
    circulo(sl, x, y, d, fundo)
    t = d * 0.58
    icone(sl, nome, x + (d - t) / 2, y + (d - t) / 2, t, cor)


def cores_badge(i: int):
    """Alterna azul e laranja entre os itens de um slide."""
    return (LAR_BG, H_LARANJA) if i % 2 == 1 else (AZUL_BG, H_AZUL)


def _texto_na_forma(s, texto, size, cor, bold=True):
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = texto
    r.font.name, r.font.size, r.font.bold = FONTE, Pt(size), bold
    r.font.color.rgb = cor
    return s


def pill(sl, rotulo, x, y, w, h, fundo=AZUL_BG, cor=AZUL, size=12):
    return _texto_na_forma(rect(sl, x, y, w, h, fill=fundo, raio=0.5), rotulo, size, cor)


def numero(sl, n, x, y, d, fundo=AZUL):
    return _texto_na_forma(circulo(sl, x, y, d, fundo), str(n), d * 28, BRANCO)


def cartao(sl, x, y, w, h, fill=SUPERF):
    return rect(sl, x, y, w, h, fill=fill, raio=min(0.12, 0.12 / max(min(w, h), 0.5)))


ALTURA_CREDITO = 0.3  # espaço do crédito da imagem, abaixo dela


def ilustracao(sl, figura, x, y, lado):
    """Imagem quadrada com cantos arredondados (recortada no centro) e o crédito embaixo, em
    letra pequena: autor e licença da foto, ou o aviso de imagem gerada por IA. O quadrado
    inteiro, com o crédito, ocupa `lado` x `lado` polegadas."""
    lado_px = 900
    imagem = figura.imagem
    menor = min(imagem.size)
    esquerda, topo = (imagem.width - menor) // 2, (imagem.height - menor) // 2
    quadrada = imagem.convert("RGB").crop((esquerda, topo, esquerda + menor, topo + menor))
    quadrada = quadrada.resize((lado_px, lado_px), Image.Resampling.LANCZOS).convert("RGBA")
    mascara = Image.new("L", (lado_px * 2, lado_px * 2), 0)
    ImageDraw.Draw(mascara).rounded_rectangle(
        (0, 0, lado_px * 2 - 1, lado_px * 2 - 1), radius=lado_px * 2 // 16, fill=255
    )
    quadrada.putalpha(mascara.resize((lado_px, lado_px), Image.Resampling.LANCZOS))
    buffer = io.BytesIO()
    quadrada.save(buffer, "PNG")
    buffer.seek(0)
    tamanho = lado - ALTURA_CREDITO
    x_imagem = x + (lado - tamanho) / 2
    sl.shapes.add_picture(buffer, I(x_imagem), I(y), I(tamanho), I(tamanho))
    txt_fit(
        sl,
        figura.credito,
        x_imagem,
        y + tamanho + 0.06,
        tamanho,
        ALTURA_CREDITO - 0.08,
        9,
        italic=True,
        cor=CINZA,
        align="r",
    )


# Transição e animação

_NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _transicao(sl) -> None:
    sl._element.append(
        etree.fromstring(f'<p:transition xmlns:p="{_NS_P}" spd="med"><p:fade/></p:transition>')
    )


def _animar_entrada(sl, ids) -> None:
    """Gráficos surgem com fade logo depois da transição do slide."""
    if not ids:
        return
    efeitos, n = [], 5
    for i, spid in enumerate(ids):
        tipo = "afterEffect" if i == 0 else "withEffect"
        efeitos.append(
            f'<p:par><p:cTn id="{n}" presetID="10" presetClass="entr" presetSubtype="0" '
            f'fill="hold" grpId="0" nodeType="{tipo}"><p:stCondLst><p:cond delay="{i * 150}"/>'
            f"</p:stCondLst><p:childTnLst>"
            f'<p:set><p:cBhvr><p:cTn id="{n + 1}" dur="1" fill="hold"><p:stCondLst>'
            f'<p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/>'
            f"</p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName>"
            f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{n + 2}" '
            f'dur="700"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
            f"</p:childTnLst></p:cTn></p:par>"
        )
        n += 3
    sl._element.append(
        etree.fromstring(
            f'<p:timing xmlns:p="{_NS_P}"><p:tnLst><p:par>'
            f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" '
            f'nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst>'
            f'<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/>'
            f'</p:cond></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold">'
            f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(efeitos)}'
            f"</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst>"
            f'</p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/>'
            f'</p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" '
            f'delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
            f"</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"
        )
    )


# Montagem dos slides


@dataclass
class Ctx:
    prs: Presentation
    total: int
    atual: int = 0
    agenda: list = field(default_factory=list)
    graficos: list = field(default_factory=list)  # ids a animar no slide atual


def nova_apresentacao(total: int) -> Ctx:
    registrar_fontes()
    prs = Presentation()
    prs.slide_width, prs.slide_height = I(W), I(H)
    return Ctx(prs, total)


def novo_slide(ctx: Ctx, notas: str = ""):
    ctx.atual += 1
    ctx.graficos = []
    sl = ctx.prs.slides.add_slide(ctx.prs.slide_layouts[6])
    set_bg(sl, BRANCO)
    if notas:
        sl.notes_slide.notes_text_frame.text = notas
    return sl


def finalizar(ctx: Ctx, sl) -> None:
    _transicao(sl)
    _animar_entrada(sl, ctx.graficos)


def pagina(ctx: Ctx, sl, cor=AZUL) -> None:
    txt(
        sl,
        f"Página {ctx.atual} de {ctx.total}",
        X1 - 3,
        6.95,
        3,
        0.3,
        size=10.5,
        italic=True,
        cor=cor,
        align="r",
    )


def barra_lateral(sl) -> None:
    """Faixa vertical do modelo: topo laranja, corte diagonal cinza, corpo azul."""
    b = 0.5
    poligono(sl, [(0, 0), (b, 0), (b, 0.85), (0, 1.2)], LARANJA)
    poligono(sl, [(0, 1.2), (b, 0.85), (b, 1.45), (0, 1.8)], rgb("#CDCDCD"))
    poligono(sl, [(0, 1.8), (b, 1.45), (b, H), (0, H)], AZUL)


def moldura_interna(ctx: Ctx, sl) -> None:
    """Barra lateral, logo SENAI no alto à direita e número da página."""
    barra_lateral(sl)
    sl.shapes.add_picture(str(LOGO_SENAI), I(X1 - 1.6), I(0.62), I(1.6))
    pagina(ctx, sl)


def cabecalho(ctx: Ctx, sl, titulo: str, subtitulo: str = "") -> None:
    moldura_interna(ctx, sl)
    txt_fit(sl, titulo, X0, 0.55, 9.6, 0.75, 32, bold=True, anchor="m")
    if subtitulo:
        txt_fit(sl, subtitulo, X0, 1.3, 9.6, 0.42, 15, cor=CINZA)


def faixa_destaque(sl, texto, y, h=0.7, icone_nome="lightbulb") -> None:
    rect(sl, X0, y, X1 - X0, h, fill=AZUL, raio=0.18)
    icone(sl, icone_nome, X0 + 0.3, y + (h - 0.38) / 2, 0.38, H_BRANCO)
    txt_fit(
        sl,
        texto,
        X0 + 0.85,
        y + 0.08,
        X1 - X0 - 1.1,
        h - 0.16,
        15,
        bold=True,
        cor=BRANCO,
        anchor="m",
    )


def grafico(ctx: Ctx, sl, fn, x, y, w, h, *args):
    """Desenha o gráfico no tamanho físico exato da área, sem distorção."""
    pic = sl.shapes.add_picture(fn(w, h, *args), I(x), I(y), I(w), I(h))
    ctx.graficos.append(pic.shape_id)
    return pic


# Gráficos (matplotlib, Open Sans, HD)


def _estilo_graficos() -> None:
    matplotlib.rcParams.update(
        {
            "font.family": FONTE,
            "font.size": 11,
            "axes.labelsize": 11,
            "xtick.labelsize": 11,
            "ytick.labelsize": 10.5,
            "axes.labelcolor": H_CINZA,
            "xtick.color": H_CINZA,
            "ytick.color": H_CINZA,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.spines.left": False,
            "axes.edgecolor": H_LINHA,
            "axes.linewidth": 1.0,
            "axes.facecolor": H_BRANCO,
            "figure.facecolor": H_BRANCO,
            "grid.color": "#EDF0F6",
            "grid.linewidth": 0.9,
            "xtick.major.size": 0,
            "ytick.major.size": 0,
            "legend.frameon": False,
            "legend.fontsize": 10.5,
            "lines.solid_capstyle": "round",
        }
    )


def figura(w, h):
    registrar_fontes()
    _estilo_graficos()
    fig = Figure(figsize=(w, h), layout="constrained")
    return fig, fig.subplots()


def png(fig) -> io.BytesIO:
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", dpi=DPI, facecolor=fig.get_facecolor())
    buffer.seek(0)
    return buffer


def suave(x, y, pontos=240):
    """Interpolação cúbica monótona (Fritsch-Carlson): curvas fluidas que não passam dos
    dados reais. Ignora os trechos com NaN."""
    x, y = np.asarray(x, float), np.asarray(y, float)
    ok = ~np.isnan(y)
    x, y = x[ok], y[ok]
    if len(x) < 3:
        return x, y
    dx, dy = np.diff(x), np.diff(y)
    m = dy / dx
    t = np.zeros_like(y)
    t[0], t[-1] = m[0], m[-1]
    for i in range(1, len(x) - 1):
        if m[i - 1] * m[i] > 0:
            t[i] = (
                3
                * (dx[i - 1] + dx[i])
                / ((2 * dx[i] + dx[i - 1]) / m[i - 1] + (dx[i] + 2 * dx[i - 1]) / m[i])
            )
    xs = np.linspace(x[0], x[-1], pontos)
    k = np.clip(np.searchsorted(x, xs) - 1, 0, len(x) - 2)
    hh = dx[k]
    s = (xs - x[k]) / hh
    h00, h10 = 2 * s**3 - 3 * s**2 + 1, s**3 - 2 * s**2 + s
    h01, h11 = -2 * s**3 + 3 * s**2, s**3 - s**2
    return xs, h00 * y[k] + h10 * hh * t[k] + h01 * y[k + 1] + h11 * hh * t[k + 1]


# Slides fixos: capa, agenda, divisória e encerramento


def slide_capa(ctx: Ctx, deck: dict, imagem=None) -> None:
    sl = novo_slide(ctx, deck.get("notas_capa") or "")
    com_logo_curso = deck.get("logo_curso", False)
    if com_logo_curso:
        sl.shapes.add_picture(logo_curso(), I(0.75), I(0.6), height=I(1.7))
        sl.shapes.add_picture(str(LOGO_SENAI), I(W - 0.75 - 2.2), I(0.75), I(2.2))
    else:
        sl.shapes.add_picture(str(LOGO_SENAI), I(0.75), I(0.75), I(2.6))
    largura = 11.8
    if imagem is not None:
        lado = 3.9
        ilustracao(sl, imagem, W - 0.75 - lado, 1.85, lado)
        largura = W - 0.75 - lado - 0.4 - 0.75

    rect(sl, 0.75, 3.0, 1.1, 0.08, fill=LARANJA)
    txt_fit(sl, deck["titulo"], 0.75, 3.2, largura, 1.85, 50, bold=True, espaco=0.95)
    linha = f"UC: {deck['uc']}" if deck.get("uc") else deck.get("subtitulo") or ""
    if linha:
        txt_fit(sl, linha, 0.75, 5.12, largura, 0.55, 22, bold=True, italic=True)
    x = 0.75
    for tag in (deck.get("tags") or [])[:5]:
        larg = min(3.2, 0.4 + 0.1 * len(tag))
        if x + larg > 0.75 + largura:
            break
        pill(sl, tag, x, 5.85, larg, 0.38)
        x += larg + 0.15
    rect(sl, 0, H - 0.12, W, 0.12, fill=AZUL)
    rect(sl, 0, H - 0.12, 2.6, 0.12, fill=LARANJA)
    finalizar(ctx, sl)


def slide_agenda(ctx: Ctx) -> None:
    sl = novo_slide(ctx, "Roteiro da apresentação.")
    cabecalho(ctx, sl, "Agenda", "O que vamos ver hoje")
    itens = ctx.agenda
    colunas = 1 if len(itens) <= 5 else 2
    por_coluna = -(-len(itens) // colunas)
    cw = (X1 - X0 - 0.6 * (colunas - 1)) / colunas
    rh = min(0.82, (Y1 - Y0) / por_coluna)
    for i, titulo in enumerate(itens):
        c, r = divmod(i, por_coluna)
        x, y = X0 + c * (cw + 0.6), Y0 + r * rh
        txt(sl, f"{i + 1:02d}", x, y, 0.8, rh - 0.12, size=24, bold=True, cor=LARANJA, anchor="m")
        txt_fit(sl, titulo, x + 0.9, y, cw - 0.9, rh - 0.12, 17, bold=True, anchor="m")
        rect(sl, x, y + rh - 0.06, cw, 0.012, fill=LINHA)
    finalizar(ctx, sl)


def slide_divisoria(ctx: Ctx, n: int, titulo: str, descricao: str) -> None:
    sl = novo_slide(ctx, f"Início da seção {n}: {titulo}.")
    sl.shapes.add_picture(fundo_azul(), 0, 0, I(W), I(H))
    rect(sl, 1.27, 1.35, 1.0, 0.09, fill=LARANJA)
    txt(sl, "SEÇÃO", 1.27, 1.55, 4, 0.4, size=16, bold=True, cor=AZUL_TXT)
    txt(sl, f"{n:02d}", 1.15, 1.85, 6, 2.3, size=120, bold=True, cor=BRANCO, anchor="m")
    txt_fit(sl, titulo, 1.27, 4.35, 10.5, 0.85, 44, bold=True, cor=BRANCO)
    if descricao:
        txt_fit(sl, descricao, 1.27, 5.25, 10.5, 0.5, 18, cor=AZUL_TXT)
    sl.shapes.add_picture(logo_branco(), I(X1 - 1.9), I(0.7), I(1.9))
    pagina(ctx, sl, BRANCO)
    finalizar(ctx, sl)


CONTATOS = [
    (
        "location_on",
        "Rodovia Admar Gonzaga, 2765 | Itacorubi\n88034-001 | Florianópolis - SC",
        False,
    ),
    ("language", "sc.senai.br", True),
    ("mail", "faleconosco@fiesc.com.br", False),
    ("call", "(48) 3231 4100  ·  0800 048 1212", False),
]


def slide_encerramento(ctx: Ctx) -> None:
    sl = novo_slide(ctx, "Encerramento e contatos.")
    sl.shapes.add_picture(fundo_azul(), 0, 0, I(W), I(H))
    sl.shapes.add_picture(logo_branco(), I(1.3), I(3.05), I(4.6))
    rect(sl, W / 2, 1.95, 0.02, 3.6, fill=BRANCO)
    cx = W / 2 + 0.55
    txt(sl, "Obrigado!", cx, 1.75, 5.5, 0.8, size=40, bold=True, cor=BRANCO)
    y = 2.8
    for nome, texto, destaque in CONTATOS:
        linhas = texto.count("\n") + 1
        icone(sl, nome, cx, y + 0.02, 0.32, H_LARANJA if destaque else H_BRANCO)
        txt(
            sl,
            texto,
            cx + 0.5,
            y,
            5.3,
            0.33 * linhas,
            size=14,
            bold=destaque,
            italic=destaque,
            cor=BRANCO,
            espaco=1.15,
        )
        y += 0.33 * linhas + 0.22
    pagina(ctx, sl, BRANCO)
    finalizar(ctx, sl)
