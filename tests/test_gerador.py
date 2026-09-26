import copy

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from senai_slides import diagrama, previa, tema
from senai_slides.gerador import gerar, roteiro
from senai_slides.imagens import nome_de_pasta
from senai_slides.layouts import LAYOUTS
from senai_slides.validador import validar_deck


def _item(n, **campos):
    return {
        "icone": "lightbulb",
        "titulo": f"Item {n}",
        "texto": f"Texto {n}.",
        "valor": f"{n}0%",
        **campos,
    }


def _itens(n):
    return [_item(i) for i in range(1, n + 1)]


# Um slide de cada layout, com o mínimo que cada um pede.
SLIDES = {
    "topicos": {"itens": _itens(3), "destaque": "Resumo do slide."},
    "cartoes": {"itens": _itens(3)},
    "numeros": {"itens": _itens(3)},
    "processo": {"itens": _itens(4)},
    "linha_do_tempo": {"itens": _itens(3)},
    "comparacao": {"colunas": [{"titulo": "Antes", "itens": ["a", "b", "c"]},
                               {"titulo": "Depois", "itens": ["d", "e", "f"]}]},
    "grafico": {"grafico": {"tipo": "linhas", "categorias": ["2023", "2024", "2025"],
                            "series": [{"nome": "Turma A", "valores": [10, 12.5, 14]}],
                            "fonte": "Dados ilustrativos"}, "itens": [_item(1)]},
    "destaque": {"destaque": "Uma frase de impacto."},
    "tabela": {
        "tabela": {"cabecalho": ["Curso", "Horas"], "linhas": [["Básico", 40], ["SEP", 40]]}
    },
    "ilustracao": {"itens": _itens(3), "imagem": {"foto": "a robot reading a book"}},
    "diagrama": {"itens": _itens(6)},
}  # fmt: skip


def _deck(**campos):
    slides = [{"layout": layout, "titulo": f"Slide {layout}", "notas": f"Notas {layout}.", **resto}
              for layout, resto in SLIDES.items()]  # fmt: skip
    return {
        "titulo": "Aula de Teste",
        "imagem": "a classroom",
        "slides": copy.deepcopy(slides),
        **campos,
    }


def _textos(slide) -> str:
    return " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def _figuras(slide) -> list:
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_o_deck_de_teste_passa_no_validador():
    assert validar_deck(_deck())[0] == []


def test_todo_layout_do_validador_tem_desenho():
    assert set(SLIDES) == set(LAYOUTS)


# Roteiro


def test_roteiro_com_agenda_e_divisorias():
    deck = _deck(secoes=[{"nome": "A", "descricao": "Primeira"}, {"nome": "B"}])
    for i, slide in enumerate(deck["slides"]):
        slide["secao"] = "A" if i < 5 else "B"
    tipos = [tipo for tipo, _ in roteiro(deck)]
    assert tipos[:3] == ["capa", "agenda", "divisoria"]
    assert tipos.count("divisoria") == 2
    assert tipos.count("conteudo") == len(SLIDES)
    assert tipos[-1] == "encerramento"


def test_sem_agenda_com_poucos_slides():
    deck = _deck()
    deck["slides"] = deck["slides"][:3]
    assert [tipo for tipo, _ in roteiro(deck)] == [
        "capa",
        "conteudo",
        "conteudo",
        "conteudo",
        "encerramento",
    ]


def test_nome_da_pasta():
    assert nome_de_pasta("Aula 1: Elétrica Básica!") == "aula_1_eletrica_basica"
    with pytest.raises(ValueError, match="letras ou números"):
        nome_de_pasta("!!!")


# Apresentação gerada


@pytest.fixture(scope="module")
def gerada(tmp_path_factory):
    """Gera o deck com todos os layouts uma vez só (com a Bedrock e o Commons falsos)."""
    from senai_slides import imagens
    from tests.conftest import BedrockFalsa, CommonsFalso

    falsa, commons = BedrockFalsa(), CommonsFalso()
    originais = imagens._bedrock, imagens._api_commons, imagens._abrir, tema.DPI
    imagens._bedrock, imagens._api_commons, imagens._abrir = (
        (lambda: falsa),
        commons.api,
        commons.abrir,
    )
    tema.DPI = 60
    try:
        pasta = gerar(_deck(), tmp_path_factory.mktemp("saida"))
    finally:
        imagens._bedrock, imagens._api_commons, imagens._abrir, tema.DPI = originais
    return pasta, Presentation(pasta / "aula_de_teste.pptx"), (falsa, commons)


def test_gera_pptx_previas_e_pdf(gerada):
    pasta, prs, _ = gerada
    total = 1 + 1 + len(SLIDES) + 1  # capa, agenda, conteúdo, encerramento
    assert len(prs.slides) == total
    assert len(list(pasta.glob("slide_*.png"))) == total
    assert (pasta / "aula_de_teste.pdf").read_bytes().startswith(b"%PDF")
    assert Image.open(pasta / "slide_01.png").size == (1920, 1080)


def test_cada_slide_de_conteudo_tem_titulo_notas_e_pagina(gerada):
    _, prs, _ = gerada
    conteudo = list(prs.slides)[2:-1]
    for layout, slide in zip(SLIDES, conteudo, strict=True):
        textos = _textos(slide)
        esperado = f"Slide {layout}".upper() if layout == "diagrama" else f"Slide {layout}"
        assert esperado in textos
        assert f"de {len(prs.slides)}" in textos
        assert slide.notes_slide.notes_text_frame.text == f"Notas {layout}."


def test_imagens_da_capa_e_da_ilustracao(gerada):
    pasta, prs, (bedrock, commons) = gerada
    # A capa pede uma imagem criada (texto sozinho); a ilustração, uma foto real.
    assert [c["prompt"].split(".")[0] for c in bedrock.chamadas] == ["a classroom"]
    assert commons.buscas == ["a robot reading a book filetype:bitmap"]
    assert len(list((pasta / "imagens").glob("*.jpg"))) == 2
    slides = list(prs.slides)
    assert "Imagem gerada por IA" in _textos(slides[0])
    ilustracao = slides[2 + list(SLIDES).index("ilustracao")]
    assert "Foto: Fulano, CC BY-SA 4.0, via Wikimedia Commons" in _textos(ilustracao)
    grandes = [f for f in _figuras(ilustracao) if f.width > tema.I(3)]
    assert len(grandes) == 1  # a ilustração; os outros são ícones e logo


def test_agenda_lista_os_titulos(gerada):
    _, prs, _ = gerada
    agenda = _textos(list(prs.slides)[1])
    assert "Agenda" in agenda
    assert all(f"Slide {layout}" in agenda for layout in SLIDES)


def test_sem_imagens_nao_busca_nem_gera(tmp_path, bedrock, commons):
    pasta = gerar(_deck(), tmp_path, imagens=False)
    prs = Presentation(pasta / "aula_de_teste.pptx")
    assert bedrock.chamadas == [] and commons.buscas == []
    ilustracao = list(prs.slides)[2 + list(SLIDES).index("ilustracao")]
    assert not [f for f in _figuras(ilustracao) if f.width > tema.I(3)]
    assert "Item 1" in _textos(ilustracao)  # vira um slide de tópicos


def test_logo_do_curso_na_capa(tmp_path):
    deck = _deck(logo_curso=True)
    deck["slides"] = deck["slides"][:1]
    pasta = gerar(deck, tmp_path, imagens=False)
    capa = Presentation(pasta / "aula_de_teste.pptx").slides[0]
    assert len(_figuras(capa)) == 2  # logo "Técnico DESI" e logo SENAI


def test_gerar_de_novo_apaga_previas_antigas(tmp_path):
    deck = _deck()
    gerar(deck, tmp_path, imagens=False)
    deck["slides"] = deck["slides"][:1]
    pasta = gerar(deck, tmp_path, imagens=False)
    assert [p.name for p in sorted(pasta.glob("slide_*.png"))] == [
        "slide_01.png",
        "slide_02.png",
        "slide_03.png",
    ]


# Texto que se ajusta à caixa


def test_texto_curto_fica_no_tamanho_maximo():
    assert tema.tamanho_que_cabe("Curto", 4, 1, 20) == 20


def test_texto_longo_diminui_ate_caber():
    texto = "Palavra " * 12
    tamanho = tema.tamanho_que_cabe(texto, 3, 0.8, 20)
    assert tema.TAMANHO_MINIMO < tamanho < 20


def test_palavra_longa_sozinha_tambem_diminui():
    # O fit_text do python-pptx falhava aqui e deixava a fonte no tamanho máximo.
    tamanho = tema.tamanho_que_cabe("DESENERGIZAÇÃO", 320 / 144, 1, 26, bold=True)
    assert tamanho < 26


# Diagrama e prévia


def test_fundo_do_diagrama_e_transparente_e_do_tamanho_do_slide():
    fundo = diagrama.fundo(6)
    assert fundo.size == (1920 * diagrama.ESCALA, 1080 * diagrama.ESCALA)
    assert fundo.mode == "RGBA"
    assert fundo.getpixel((10, 10))[3] == 0


@pytest.mark.parametrize("n", range(3, 7))
def test_cores_do_diagrama_vao_do_laranja_ao_azul(n):
    cores = diagrama.cores(n)
    assert (cores[0], cores[-1]) == (tema.H_LARANJA, tema.H_AZUL)
    assert len(diagrama.posicoes(n)) == n


def test_previa_desenha_texto_e_formas(gerada):
    _, prs, _ = gerada
    imagem = previa.renderizar(prs, prs.slides[0], 960)
    assert imagem.size == (960, 540)
    assert len(imagem.getcolors(100_000)) > 10
